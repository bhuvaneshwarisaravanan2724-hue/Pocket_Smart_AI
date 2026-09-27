from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION


def create_budget_report(
    income: float,
    total_expenses: float,
    recommendation: str,
    expenses,
    output_path: str
):
    document = Document()

    # Page margins
    section = document.sections[0]
    section.top_margin = Pt(50)
    section.bottom_margin = Pt(50)
    section.left_margin = Pt(60)
    section.right_margin = Pt(60)

    # Title
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run("POCKET SMART AI")
    run.bold = True
    run.font.size = Pt(22)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitle.add_run(
        "SMART BUDGET & RECOMMENDATION REPORT"
    )
    run.bold = True
    run.font.size = Pt(14)

    document.add_paragraph()

    # 1. User Budget Details
    document.add_heading("1. User Budget Details", level=1)

    budget_table = document.add_table(rows=2, cols=2)
    budget_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    budget_table.style = "Table Grid"

    budget_table.cell(0, 0).text = "Monthly Income"
    budget_table.cell(0, 1).text = f"₹{income:,.2f}"

    budget_table.cell(1, 0).text = "Report Type"
    budget_table.cell(1, 1).text = "AI-Powered Budget Report"

    # 2. Expense Summary
    document.add_heading("2. Expense Summary", level=1)

    expense_table = document.add_table(rows=1, cols=3)
    expense_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    expense_table.style = "Table Grid"

    expense_table.cell(0, 0).text = "Category"
    expense_table.cell(0, 1).text = "Description"
    expense_table.cell(0, 2).text = "Amount"

    for expense in expenses:
        row = expense_table.add_row().cells

        row[0].text = expense.category
        row[1].text = expense.description or "-"
        row[2].text = f"₹{expense.amount:,.2f}"

    total_row = expense_table.add_row().cells

    total_row[0].text = "TOTAL"
    total_row[1].text = ""
    total_row[2].text = f"₹{total_expenses:,.2f}"

    # 3. Remaining Budget
    remaining = income - total_expenses

    document.add_heading("3. Remaining Budget", level=1)

    remaining_table = document.add_table(rows=1, cols=2)
    remaining_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remaining_table.style = "Table Grid"

    remaining_table.cell(0, 0).text = "Remaining Amount"
    remaining_table.cell(0, 1).text = f"₹{remaining:,.2f}"

    # 4. AI Recommendation
    document.add_heading("4. AI Recommendation", level=1)

    document.add_paragraph(
        recommendation
    )

    # 5. Conclusion
    document.add_heading("5. Conclusion", level=1)

    document.add_paragraph(
        "Pocket Smart AI provides a personalized budget summary "
        "and AI-powered recommendation to help users manage "
        "their spending and plan their savings."
    )

    # Save document
    document.save(output_path)